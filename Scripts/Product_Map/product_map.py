import openpyxl
import networkx
import os
import matplotlib.pyplot as plt
import re
from pyvis.network import Network
import plotly.graph_objects as go




def import_BOMs_net():
    script_path = os.path.dirname(__file__)
    
    graph = networkx.Graph()

    wb = openpyxl.load_workbook(script_path + "/BOM_GL_JDE.xlsx")
    sheet = wb.active

    children = [re.sub(r'\s+', '_', i.value) for i in sheet['H'][1:]]
    parents = [re.sub(r'\s+', '_', i.value) for i in sheet['K'][1:]]

    data = zip(children, parents)

    for i in data:
        graph.add_edge(i[0],i[1])


    nt = Network(height="1000px", width="2000px", notebook=True)
    nt.from_nx(graph)
    nt.barnes_hut()
    nt.show("nx.html")
    # networkx.draw(graph)
    # plt.show()

    # [print(x[0].value+","+ x[1].value) for x in data]

def load_routings():
    dictionary = {}
    script_path = os.path.dirname(__file__)
    wb = openpyxl.load_workbook(script_path + "/JDE_routings.xlsx")
    routings = wb["Sheet1"]
    Work_center = routings["L"][1:]
    Product = routings["I"][1:]
    
def import_item_dictionary():
    script_path = os.path.dirname(__file__)
    wb = openpyxl.load_workbook(script_path + "/2nd_item_no_register.xlsx")
    sheet = wb["Sheet1"]
    item_no = [re.sub(r'\s+', '_', i.value) for i in sheet["A"][1:]]
    descrip = [re.sub(r'\s+','_', i.value) for i in sheet["C"][1:]]
    dictionary = dict(zip(item_no, descrip))
    return dictionary
    

def import_BOMs_Sankey():
    item_dict = import_item_dictionary()
    
    script_path = os.path.dirname(__file__)
    
    wb = openpyxl.load_workbook(script_path + "/BOM_GL_JDE.xlsx")
    sheet = wb['Sheet1']

    children = [re.sub(r'\s+', '_', i.value) for i in sheet['H'][1:]]
    parents = [re.sub(r'\s+', '_', i.value) for i in sheet['K'][1:]]

    # These lines are to alter the visibility of the packaging products.
    pairs = zip(parents, children)
    pairs = [ i if i[0][0] != "2" else i[::-1] for i in pairs] # This places packaging materials at the end of the graph
    # pairs = [i for i in pairs if i[0][0] != "2"] # This removes packaging materials all together
    parents, children = zip(*pairs) # this unzips the alterations

    
    ordered_set = list(set(parents+children))

    rmats = {i: {"color":"red", "X_pos":0.1} for i in ordered_set if re.match(r"\ARM.*", i) }
    drier_product = {i: {"color":"purple", "X_pos":0.3} for i in ordered_set if re.match(r'\A.*DP_WIP_GL', i) }
    milled_product = {i: {"color":"lightskyblue", "X_pos":0.4} for i in ordered_set if re.match(r'\A.*+M\d+_WIP_GL',i)}
    bentonorit = {i: {"color":"black", "X_pos":0.5} for i in ordered_set if re.match(r"\ABN.*+[^R]_[^W][^I][^P]_GL", i)}
    finished_goods = {i: {"color":"green", "X_pos":0.6} for i in ordered_set if re.match(r"\A[^B][^N].*_[^W][^I][^P]_GL", i)}
    packaging = {item_dict[i]:{"color":"blue", "X_pos":0.9} for i in ordered_set if i[0]=="2"}
    r_plant = {i: {"color":"pink", "X_pos":0.7} for i in ordered_set if re.match(r"\A.*R_WIP_GL",i)}
    # packaging = {item_dict[i]: {"color":"blue", "X_pos":0.9} for i in packaging}

    list_of_dicts = [rmats, drier_product, milled_product, bentonorit, finished_goods, packaging, r_plant]

    options = {}
    for d in list_of_dicts:
        options.update(d)



    # These rename packaging materails were information is available
    ordered_set = [item_dict[i] if i[0]=="2" and i in item_dict else i for i in ordered_set]
    children = [item_dict[i] if i[0]=="2" and i in item_dict else i for i in children]
    parents = [item_dict[i] if i[0]=="2" and i in item_dict else i for i in parents]


    fig  = go.Figure(go.Sankey(
        node = dict(
            pad = 5,
            thickness = 10,
            line = dict(color = "black",width = 0.5),
            label = ordered_set,
            x = [options[i]["X_pos"]if i in options else 1 for i in ordered_set],
            color = [options[i]["color"] if i in options else "blue" for i in ordered_set],
            pad=5,
            ),
        link =dict(
            source = [ordered_set.index(i) for i in parents],
            target = [ordered_set.index(i) for i in children],
            value = [1 for i in parents],
            hovercolor = "gold"
            )
        )
    )
    fig.write_html('BOM_Sankey_Diagram.html')
    fig.show()


# import_BOMs_net()
import_BOMs_Sankey()
